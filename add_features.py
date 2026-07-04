from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

wb = load_workbook('Business_Sales_Expense_Tracker_v2.xlsx')

NAVY = "1F3864"
LIGHT_BLUE = "D9E2F3"
WHITE_FONT = Font(color="FFFFFF", bold=True, name="Arial", size=11)
HEADER_FILL = PatternFill("solid", start_color=NAVY)
TITLE_FONT = Font(bold=True, size=16, color=NAVY, name="Arial")
SUBTITLE_FONT = Font(italic=True, size=10, color="595959", name="Arial")
LABEL_FONT = Font(bold=True, name="Arial", size=11)
NORMAL_FONT = Font(name="Arial", size=11)
BLUE_INPUT = Font(color="0000FF", name="Arial", size=11)
BLACK_FORMULA = Font(color="000000", name="Arial", size=11)
thin = Side(style="thin", color="B7B7B7")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# ============ NEW SHEET: PRODUCT LIST ============
ws_p = wb.create_sheet("Product List", 1)  # place right after Instructions
ws_p.sheet_view.showGridLines = False
ws_p.column_dimensions['A'].width = 28
ws_p.column_dimensions['B'].width = 20

ws_p['A1'] = "Product Name"
ws_p['B1'] = "Default Price (₹)"
for c in ['A1', 'B1']:
    ws_p[c].font = WHITE_FONT
    ws_p[c].fill = HEADER_FILL
    ws_p[c].alignment = Alignment(horizontal="center", vertical="center")
    ws_p[c].border = BORDER
ws_p.row_dimensions[1].height = 26

sample_products = [
    ("Notebook Set", 150),
    ("Water Bottle", 250),
    ("Backpack", 900),
    ("Pen Pack", 60),
]
for i, (name, price) in enumerate(sample_products, start=2):
    ws_p[f'A{i}'] = name
    ws_p[f'A{i}'].font = BLUE_INPUT
    ws_p[f'A{i}'].border = BORDER
    ws_p[f'B{i}'] = price
    ws_p[f'B{i}'].font = BLUE_INPUT
    ws_p[f'B{i}'].number_format = '₹#,##0'
    ws_p[f'B{i}'].border = BORDER

# extend blank rows for more products (up to row 31 = 30 products)
for i in range(2 + len(sample_products), 32):
    ws_p[f'A{i}'].font = BLUE_INPUT
    ws_p[f'A{i}'].border = BORDER
    ws_p[f'B{i}'].font = BLUE_INPUT
    ws_p[f'B{i}'].number_format = '₹#,##0'
    ws_p[f'B{i}'].border = BORDER

ws_p['D1'] = "Add your product names & prices here."
ws_p['D1'].font = SUBTITLE_FONT
ws_p['D2'] = "This list feeds the dropdown in Sales Entry."
ws_p['D2'].font = SUBTITLE_FONT

# ============ SALES ENTRY: add dropdown on Product Name ============
ws1 = wb['Sales Entry']
dv_product = DataValidation(type="list", formula1="='Product List'!$A$2:$A$31", allow_blank=True, showDropDown=False)
ws1.add_data_validation(dv_product)
for row in range(2, 501):
    dv_product.add(ws1[f'B{row}'])

# ============ DASHBOARD: add "Today's Sales" snapshot ============
ws3 = wb['Dashboard']
ws3.column_dimensions['G'].width = 4
ws3.column_dimensions['H'].width = 22
ws3.column_dimensions['I'].width = 16
ws3.column_dimensions['J'].width = 18

ws3['H2'] = "Today's Sales Snapshot"
ws3['H2'].font = LABEL_FONT
ws3['H3'] = "=\"Auto-updates for \"&TEXT(TODAY(),\"DD-MMM-YYYY\")"
ws3['H3'].font = SUBTITLE_FONT

ws3['H5'] = "Product"
ws3['I5'] = "Qty Sold Today"
ws3['J5'] = "Revenue Today (₹)"
for c in ['H5', 'I5', 'J5']:
    ws3[c].font = WHITE_FONT
    ws3[c].fill = HEADER_FILL
    ws3[c].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws3[c].border = BORDER
ws3.row_dimensions[5].height = 28

for i in range(30):
    r = 6 + i
    prow = 2 + i
    ws3[f'H{r}'] = f"='Product List'!A{prow}"
    ws3[f'H{r}'].font = BLACK_FORMULA
    ws3[f'H{r}'].border = BORDER
    ws3[f'I{r}'] = f'=IF(H{r}="","",SUMIFS(\'Sales Entry\'!$C$2:$C$500,\'Sales Entry\'!$B$2:$B$500,H{r},\'Sales Entry\'!$A$2:$A$500,TODAY()))'
    ws3[f'I{r}'].font = BLACK_FORMULA
    ws3[f'I{r}'].border = BORDER
    ws3[f'J{r}'] = f'=IF(H{r}="","",SUMIFS(\'Sales Entry\'!$E$2:$E$500,\'Sales Entry\'!$B$2:$B$500,H{r},\'Sales Entry\'!$A$2:$A$500,TODAY()))'
    ws3[f'J{r}'].font = BLACK_FORMULA
    ws3[f'J{r}'].number_format = '₹#,##0'
    ws3[f'J{r}'].border = BORDER

ws3['H37'] = "Total Today"
ws3['H37'].font = LABEL_FONT
ws3['H37'].fill = PatternFill("solid", start_color=LIGHT_BLUE)
ws3['I37'] = '=SUMIFS(\'Sales Entry\'!$C$2:$C$500,\'Sales Entry\'!$A$2:$A$500,TODAY())'
ws3['I37'].font = Font(bold=True, name="Arial", size=11, color=NAVY)
ws3['I37'].border = BORDER
ws3['J37'] = '=SUMIFS(\'Sales Entry\'!$E$2:$E$500,\'Sales Entry\'!$A$2:$A$500,TODAY())'
ws3['J37'].font = Font(bold=True, name="Arial", size=11, color=NAVY)
ws3['J37'].number_format = '₹#,##0'
ws3['J37'].border = BORDER

# ============ DASHBOARD: add "All-Time Product Totals" (kitna total bika) ============
ws3['H40'] = "All-Time Product Totals"
ws3['H40'].font = LABEL_FONT
ws3['H41'] = "Product"
ws3['I41'] = "Total Qty Sold"
ws3['J41'] = "Total Revenue (₹)"
for c in ['H41', 'I41', 'J41']:
    ws3[c].font = WHITE_FONT
    ws3[c].fill = HEADER_FILL
    ws3[c].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws3[c].border = BORDER
ws3.row_dimensions[41].height = 28

for i in range(30):
    r = 42 + i
    prow = 2 + i
    ws3[f'H{r}'] = f"='Product List'!A{prow}"
    ws3[f'H{r}'].font = BLACK_FORMULA
    ws3[f'H{r}'].border = BORDER
    ws3[f'I{r}'] = f'=IF(H{r}="","",SUMIF(\'Sales Entry\'!$B$2:$B$500,H{r},\'Sales Entry\'!$C$2:$C$500))'
    ws3[f'I{r}'].font = BLACK_FORMULA
    ws3[f'I{r}'].border = BORDER
    ws3[f'J{r}'] = f'=IF(H{r}="","",SUMIF(\'Sales Entry\'!$B$2:$B$500,H{r},\'Sales Entry\'!$E$2:$E$500))'
    ws3[f'J{r}'].font = BLACK_FORMULA
    ws3[f'J{r}'].number_format = '₹#,##0'
    ws3[f'J{r}'].border = BORDER

# ============ Update Instructions sheet ============
ws0 = wb['Instructions']
ws0['B17'] = "6. Add your product names & prices in the 'Product List' tab — this powers the dropdown in Sales Entry."
ws0['B17'].font = NORMAL_FONT
ws0['B18'] = "7. In Sales Entry, click the Product Name cell and pick from the dropdown instead of typing."
ws0['B18'].font = NORMAL_FONT
ws0['B19'] = "8. Check Dashboard (right side) for Today's Sales and All-Time Product Totals — updates automatically."
ws0['B19'].font = NORMAL_FONT

wb.save('Business_Sales_Expense_Tracker_v2.xlsx')
print("done")
